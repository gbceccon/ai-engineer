"""Build gabriel_ceccon_cv_us_ai_v1.docx — US-market revision of Gabriel's CV.

Keeps the strong original narrative, repositions for AI-augmented delivery roles,
promotes Python/data work, adds AI Projects / Certifications / GitHub sections.
ATS-safe: no tables, standard headings, single column.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x3B, 0x63)
GRAY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# Base style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(4)

for s in doc.sections:
    s.top_margin = Inches(0.55)
    s.bottom_margin = Inches(0.55)
    s.left_margin = Inches(0.7)
    s.right_margin = Inches(0.7)


def para(text="", size=10.5, bold=False, color=None, space_after=4, space_before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def heading(text):
    p = para(text.upper(), size=11.5, bold=True, color=ACCENT, space_before=10, space_after=3)
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F3B63")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(10.5)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    return p


def job(title_company, dates, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.tab_stops.add_tab_stop(doc.sections[0].page_width)
    r = p.add_run(title_company)
    r.bold = True
    r.font.size = Pt(10.5)
    r2 = p.add_run("\t" + dates)
    r2.font.size = Pt(10)
    r2.font.color.rgb = GRAY
    for b in bullets:
        bullet(b)


# ---------------- Header ----------------
p = para("GABRIEL BARTZ CECCON", size=17, bold=True, color=ACCENT, space_after=1)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = para(
    "Senior Technical Program Manager — AI-augmented Delivery  |  Open Finance & Fintech  |  Regulatory Incident Management",
    size=10.5, color=GRAY, space_after=1,
)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = para(
    "gbceccon@gmail.com  •  Porto Alegre, Brazil — open to remote roles with US companies (contractor/B2B)  •  "
    "linkedin.com/in/gabrielceccon44  •  github.com/<username>/ai-engineer",
    size=9.5, color=GRAY, space_after=6,
)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------------- Summary ----------------
heading("Professional Summary")
para(
    "I translate complexity in both directions. When a regulatory incident hits a fintech platform, I trace it to "
    "the root cause across multiple engineering teams, then walk executives and regulators through exactly what it "
    "means for the business, what needs to happen, and by when. For 4 years I led end-to-end regulatory incident "
    "management for Open Finance at Sicredi — one of Brazil's largest financial cooperatives — scaling the program "
    "from 20 to 200 monthly cases (10x) and authoring the formal remediation commitment submitted to Brazil's "
    "Central Bank (BACEN)."
)
para(
    "Now pairing 18 years of fintech and distributed-systems depth — from C++ firmware at HP to cross-team incident "
    "command in core banking — with hands-on AI/LLM engineering: retrieval systems, LLM evaluation in regulated "
    "domains, and AI-augmented operational workflows. My public portfolio (GitHub) documents each system I build, "
    "with benchmarks proving each version beats the last."
)

# ---------------- AI Projects ----------------
heading("Selected AI Projects (Public Portfolio)")
job(
    "OpenPilot — AI copilot for regulatory incident management  |  github.com/<username>/ai-engineer",
    "2025 – present",
    [
        "Building an LLM-powered triage and remediation-drafting assistant for Open Finance incident cases "
        "(synthetic dataset): TF-IDF/semantic search, fine-tuned BERT classifiers, RAG over incident history "
        "and BACEN regulations with citations.",
        "Designing evaluation harnesses for LLM systems in regulated domains: golden-question sets, accuracy/"
        "cost/latency tracking, regression benchmarks between model versions.",
        "Production-grade workflow: pytest + GitHub Actions CI, versioned releases (v0.1→v1.0), documented "
        "benchmarks comparing every model iteration.",
    ],
)
para(
    "(Project scope grows as my 12-month AI engineering program progresses — see repository ROADMAP.md. "
    "Replace this project block with completed milestones as they ship.)",
    size=9, italic=True, color=GRAY,
)

# ---------------- Experience ----------------
heading("Professional Experience")
job(
    "Resolution Manager (BSA Senior) — Open Finance Program Management  |  Sicredi",
    "Jan 2022 – Present",
    [
        "Own end-to-end program management for Open Finance regulatory incidents, coordinating action plans "
        "across 5+ engineering, product, and business teams simultaneously with no direct authority — driving "
        "resolution through structured escalation, weekly accountability cadences, and stakeholder alignment.",
        "Built and scaled the Open Finance incident management program from the ground up — growing from 20 to "
        "200 monthly regulatory cases (10x) while continuously redesigning workflows to match volume and "
        "regulatory complexity.",
        "Automated regulatory SLA reconciliation with Python/pandas, joining SLA data with internal ticketing "
        "systems and closing a visibility gap between reported compliance and operational reality — the workflow "
        "now serving as the data foundation for AI-assisted triage.",
        "Designed and govern a cross-organizational incident management framework spanning engineering, product, "
        "and business stakeholders, ensuring continuous compliance with Brazil's Central Bank (BACEN) requirements.",
        "Authored the formal regulatory commitment submitted to Brazil's Central Bank — incident timeline, root "
        "cause analysis, and remediation plan with committed delivery dates — fully executed on schedule.",
        "Primary bridge between engineering teams and business/product leadership, translating root cause "
        "analysis and technical risk into business impact for General Manager and Business Director-level stakeholders.",
        "Lead a team of 3 support analysts, mentoring and delegating operational ownership to build team "
        "capability and free capacity for proactive, data-driven process improvement.",
    ],
)
job(
    "Software Engineer — Core Banking Program & Production Support  |  Agibank",
    "Nov 2018 – Jan 2022",
    [
        "Coordinated cross-company teams to plan, test, and execute regulatory-driven version updates on the "
        "Matera core banking platform — a zero-rollback environment — diagnosing and resolving live production "
        "issues in real time without service interruption.",
        "Acted as Product Owner for Card Processing, bridging engineering and business requirements across "
        "current account, investments, and registration systems.",
        "Led fault investigation and cross-team problem resolution using AWS, microservices, Spring Boot, "
        "Docker, and REST APIs; managed change and release processes (ITIL).",
    ],
)
job(
    "Technical Lead — ATM Infrastructure Support Program  |  Saque e Pague",
    "Jan 2013 – Nov 2018",
    [
        "Diagnosed a systemic network failure across a fleet of 1,100 ATMs, tracing outdated software, failed "
        "patch deployments, and gateway transaction breaks to a single root cause — built the business case that "
        "drove investment in dedicated network links, gaining 35% first-call resolution and 50% terminal "
        "refresh-rate improvement fleet-wide.",
        "Led deployment automation for the full ATM software stack across 1,100 terminals and developed "
        "supporting tooling for N2 support; coordinated N1/N2/N3 and agile dev teams as technical program lead.",
    ],
)
job(
    "Software Engineer  |  Hewlett-Packard  •  Perto S.A.",
    "Oct 2006 – Dec 2012",
    [
        "Developed a NoSQL distributed database (C, C++, STL, Boost, Perl, Thrift) and firmware for PMC "
        "processors on SAS disk enclosures / fiber channel for HP EVA Storage.",
        "Developed test applications and communication-layer APIs (serial, parallel, USB) for banking devices "
        "in C/C++; supported international clients (Israel, India, Chile, United States).",
    ],
)

# ---------------- Education ----------------
heading("Education")
para("B.Eng. Control & Automation Engineering — PUCRS (Pontifícia Universidade Católica do RS), 2001–2006")

# ---------------- Certifications ----------------
heading("Certifications (in progress)")
bullet("DeepLearning.AI — Generative AI / LLM specialization courses (2025)")
bullet("AWS Certified Machine Learning – Specialty (target: Phase 6 of AI program)")
para(
    "(Update this section as certifications complete — part of the public 12-month program on GitHub.)",
    size=9, italic=True, color=GRAY,
)

# ---------------- Skills ----------------
heading("Core Competencies")
bullet("AI & Data — Python, pandas/NumPy, scikit-learn, LLM/RAG systems, evaluation harnesses, "
       "prompt engineering, PyTorch & Hugging Face (in active practice), data pipelines", bold_prefix="")
bullet("AI-augmented delivery — incident triage automation, LLM evaluation in regulated domains, "
       "AI workflow design, model risk & governance awareness (NIST AI RMF)")
bullet("Program & cross-team leadership — coordinating delivery across teams with no direct authority, "
       "structured escalation, executive communication")
bullet("Regulatory & risk — BACEN Open Finance compliance, incident command, root cause analysis, "
       "formal remediation commitments")
bullet("Technical depth — AWS, Docker, microservices, REST APIs, SQL, C/C++, ITIL change management")
bullet("Languages — Portuguese (native) • English (professional working proficiency — daily written "
       "technical output in English)")

out = "/Users/gabrielceccon/Documents/ai-engineer/gabriel_ceccon_cv_us_ai_v1.docx"
doc.save(out)
print("saved", out)
